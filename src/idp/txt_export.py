"""Export per-file text bundles with hyperlinks to images for DOCX/PDF reconstruction.

Every processed document gets its own folder inside ``<output_root>/txts/``:

    txts/
      <file name>/
        text.md      document text with hyperlinks to its images
        images/      page renders and embedded pictures of the document
        tables.json  structured markdown tables (rows/cells)
        meta.json    source info and image inventory

Tables are kept as Markdown tables in ``text.md`` (so ``pandoc text.md -o out.docx``
rebuilds them) and additionally mirrored as structured rows in ``tables.json``.
Handwritten fragments are referenced as hyperlinks to the page image, because
manuscripts are preserved as images instead of unreliable text.
"""

from __future__ import annotations

import base64
import json
import os
import re
import shutil
import zipfile
from pathlib import Path
from urllib.parse import unquote

from idp.config import settings

_PAGE_MARKER_RE = re.compile(r"<!--\s*page\s+(\d+)\s*-->")
_SLIDE_MARKER_RE = re.compile(r"<!--\s*slide\s+(\d+)\s*-->")
_FILE_MARKER_RE = re.compile(r"<!--\s*file:\s*(.+?)\s*-->")
_IMAGE_PLACEHOLDER_RE = re.compile(r"\[Image:\s*([^\]]*)\]")
_HANDWRITTEN_RE = re.compile(r"\[HANDWRITTEN:\s*([^\]]*)\]")
_MD_IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)\s]+)\)")
_DATA_URI_RE = re.compile(r"!\[([^\]]*)\]\(data:image/([A-Za-z0-9.+-]+);base64,([^)\s]+)\)")
_TABLE_SEPARATOR_RE = re.compile(r"^:?-+:?$")
_UNSAFE_NAME_RE = re.compile(r'[<>:"/\\|?*\x00-\x1f]+')
_MAX_EMBEDDED_IMAGES = 500


def _sanitize_name(name: str) -> str:
    cleaned = _UNSAFE_NAME_RE.sub("_", name).strip().strip(".")
    return cleaned or "unnamed"


def _link_or_copy(source: Path, target: Path) -> None:
    try:
        os.link(source, target)
    except OSError:
        shutil.copy2(source, target)


def _image_extension(extension: str) -> str:
    cleaned = (extension or "").lower().lstrip(".")
    if cleaned in ("jpeg", "jpg"):
        return "jpg"
    if cleaned == "svg+xml":
        return "svg"
    return cleaned or "png"


def _is_remote(source: str) -> bool:
    return source.startswith(("http://", "https://", "ftp://", "//", "data:", "mailto:"))


def _resolve_local_source(source: str, base_dirs: list[Path]) -> Path | None:
    cleaned = unquote(source.split("?", 1)[0].split("#", 1)[0]).strip()
    if not cleaned:
        return None
    candidate = Path(cleaned)
    if candidate.is_absolute() and candidate.is_file():
        return candidate
    for base in base_dirs:
        if base is None:
            continue
        resolved = base / cleaned
        if resolved.is_file():
            return resolved
    return None


def _materialize_data_uris(images_dir: Path, text: str, prefix: str, counter: int = 0) -> tuple[str, int]:
    def _replace(match: re.Match) -> str:
        nonlocal counter
        alt = match.group(1).strip()
        extension = _image_extension(match.group(2))
        try:
            data = base64.b64decode(match.group(3))
        except Exception:
            return match.group(0)
        counter += 1
        name = f"{prefix}_{counter:04d}.{extension}"
        (images_dir / name).write_bytes(data)
        return f"![{alt or 'Image'}](images/{name})"

    return _DATA_URI_RE.sub(_replace, text), counter


def _rewrite_local_image_refs(
    text: str,
    base_dirs: list[Path],
    images_dir: Path,
    prefix: str,
    counter: int = 0,
) -> tuple[str, int]:
    def _replace(match: re.Match) -> str:
        nonlocal counter
        alt = match.group(1).strip()
        source = match.group(2).strip()
        if _is_remote(source):
            return match.group(0)
        resolved = _resolve_local_source(source, base_dirs)
        if resolved is None:
            return match.group(0)
        counter += 1
        name = f"{prefix}_{counter:04d}.{_image_extension(resolved.suffix)}"
        _link_or_copy(resolved, images_dir / name)
        return f"![{alt or resolved.stem}](images/{name})"

    return _MD_IMAGE_RE.sub(_replace, text), counter


def _extract_pdf_embedded_images(pdf_path: Path, images_dir: Path) -> dict[int, list[str]]:
    result: dict[int, list[str]] = {}
    try:
        import fitz
    except ImportError:
        return result
    try:
        doc = fitz.open(pdf_path)
    except Exception:
        return result
    by_xref: dict[int, str] = {}
    try:
        for page_index in range(len(doc)):
            try:
                infos = doc.load_page(page_index).get_images(full=True)
            except Exception:
                continue
            names: list[str] = []
            for info in infos:
                xref = info[0] if info else 0
                if not xref:
                    continue
                if xref in by_xref:
                    names.append(by_xref[xref])
                    continue
                if len(by_xref) >= _MAX_EMBEDDED_IMAGES:
                    continue
                try:
                    extracted = doc.extract_image(xref)
                except Exception:
                    continue
                data = extracted.get("image")
                if not data:
                    continue
                name = f"img_{len(by_xref) + 1:04d}.{_image_extension(extracted.get('ext', 'png'))}"
                (images_dir / name).write_bytes(data)
                by_xref[xref] = name
                names.append(name)
            if names:
                result[page_index + 1] = names
    finally:
        doc.close()
    return result


def _iter_pptx_shapes(shapes):
    for shape in shapes:
        yield shape
        nested = getattr(shape, "shapes", None)
        if nested is not None:
            yield from _iter_pptx_shapes(nested)


def _extract_pptx_images(pptx_path: Path, images_dir: Path) -> dict[int, list[str]]:
    result: dict[int, list[str]] = {}
    try:
        from pptx import Presentation
    except ImportError:
        return result
    try:
        presentation = Presentation(str(pptx_path))
    except Exception:
        return result
    for slide_no, slide in enumerate(presentation.slides, 1):
        names: list[str] = []
        for shape in _iter_pptx_shapes(slide.shapes):
            try:
                image = shape.image
            except Exception:
                continue
            try:
                data = image.blob
            except Exception:
                continue
            if not data:
                continue
            try:
                extension = image.ext
            except Exception:
                extension = "png"
            name = f"slide_{slide_no:04d}_image_{len(names) + 1}.{_image_extension(extension)}"
            (images_dir / name).write_bytes(data)
            names.append(name)
        if names:
            result[slide_no] = names
    return result


def parse_markdown_tables(text: str) -> list[dict]:
    tables: list[dict] = []
    page: int | None = None
    rows: list[list[str]] = []

    def _flush() -> None:
        nonlocal rows
        if rows:
            tables.append({"page": page, "rows": rows})
            rows = []

    for line in text.splitlines():
        page_match = _PAGE_MARKER_RE.search(line)
        if page_match is not None:
            _flush()
            page = int(page_match.group(1))
            continue
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2:
            cells = [cell.strip() for cell in stripped.strip("|").split("|")]
            if all(_TABLE_SEPARATOR_RE.match(cell) for cell in cells if cell):
                continue
            rows.append(cells)
            continue
        _flush()
    _flush()
    return tables


class TxtExporter:
    """Writes one ``txts/<file name>/`` bundle per processed document."""

    def __init__(self, output_root: Path | None = None) -> None:
        self.root = (output_root or settings.output_root) / "txts"
        self._used: set[str] = set()

    def export(
        self,
        source: Path,
        file_type: str,
        text: str,
        *,
        pages: int | None = None,
        document_type: str = "other",
        annotation: str = "",
        metadata: dict | None = None,
    ) -> Path:
        bundle_dir = self._bundle_dir(source)
        images_dir = bundle_dir / "images"
        shutil.rmtree(bundle_dir, ignore_errors=True)
        images_dir.mkdir(parents=True, exist_ok=True)

        text = text or ""
        if file_type == "pdf":
            text = self._bundle_pdf(source, images_dir, text)
        elif file_type == "docx":
            text = self._bundle_docx(source, images_dir, text)
        elif file_type == "pptx":
            text = self._bundle_pptx(source, images_dir, text)
        elif file_type in ("html", "html_folder"):
            text = self._bundle_html(source, images_dir, text, folder=file_type == "html_folder")

        header = f"<!-- source: {self._relative(source)} -->"
        (bundle_dir / "text.md").write_text(f"{header}\n\n{text.strip()}\n", encoding="utf-8")
        (bundle_dir / "tables.json").write_text(
            json.dumps({"tables": parse_markdown_tables(text)}, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        self._write_meta(bundle_dir, source, file_type, pages, document_type, annotation, metadata or {}, images_dir)
        return bundle_dir

    def _bundle_dir(self, source: Path) -> Path:
        base = _sanitize_name(source.name)
        candidate = base
        counter = 2
        while candidate.lower() in self._used:
            candidate = f"{base}_{counter}"
            counter += 1
        self._used.add(candidate.lower())
        return self.root / candidate

    @staticmethod
    def _relative(source: Path) -> str:
        try:
            return str(source.resolve().relative_to(settings.input_root.resolve()))
        except Exception:
            return source.name

    def _bundle_pdf(self, pdf_path: Path, images_dir: Path, text: str) -> str:
        page_images = self._copy_page_images(pdf_path, images_dir)
        embedded = _extract_pdf_embedded_images(pdf_path, images_dir)
        if not page_images and not embedded:
            return text

        markers = [int(value) for value in _PAGE_MARKER_RE.findall(text)]
        total_pages = len(page_images)
        pages_per_marker = 1
        if markers and total_pages and len(markers) < total_pages:
            pages_per_marker = max(1, settings.vl_max_images)

        preamble: list[str] = []
        sections: list[tuple[list[str], list[int]]] = []
        current_lines: list[str] | None = None
        current_pages: list[int] = []

        for line in text.splitlines():
            marker = _PAGE_MARKER_RE.search(line)
            if marker is not None:
                if current_lines is not None:
                    sections.append((current_lines, current_pages))
                current_lines = [line]
                current_pages = self._marker_pages(int(marker.group(1)), pages_per_marker, total_pages)
                continue
            if current_lines is None:
                preamble.append(line)
            else:
                current_lines.append(line)
        if current_lines is not None:
            sections.append((current_lines, current_pages))

        consumed: set[str] = set()
        covered: set[int] = set()
        out: list[str] = list(preamble)
        for lines, pages in sections:
            covered.update(pages)
            out.extend(lines[:1])
            for page in pages:
                name = page_images.get(page)
                if name is not None:
                    out.append(f"![Page {page}](images/{name})")
            if len(lines) > 1:
                out.extend(self._link_placeholders("\n".join(lines[1:]), pages, embedded, page_images, consumed).split("\n"))
            for page in pages:
                for name in embedded.get(page, []):
                    if name not in consumed:
                        consumed.add(name)
                        out.append(f"![Image on page {page}](images/{name})")

        uncovered = [page for page in sorted(page_images) if page not in covered]
        if uncovered:
            if out and out[-1].strip():
                out.append("")
            out.append("## Pages")
            for page in uncovered:
                out.append(f"![Page {page}](images/{page_images[page]})")
        return "\n".join(out)

    @staticmethod
    def _marker_pages(marker: int, pages_per_marker: int, total_pages: int) -> list[int]:
        start = (marker - 1) * pages_per_marker + 1
        end = min(marker * pages_per_marker, total_pages) if total_pages else start
        if end < start:
            end = start
        return list(range(start, end + 1))

    @staticmethod
    def _link_placeholders(
        block: str,
        pages: list[int],
        embedded: dict[int, list[str]],
        page_images: dict[int, str],
        consumed: set[str],
    ) -> str:
        def _image_sub(match: re.Match) -> str:
            description = match.group(1).strip() or "Image"
            for page in pages:
                for name in embedded.get(page, []):
                    if name not in consumed:
                        consumed.add(name)
                        return f"![{description}](images/{name})"
            return match.group(0)

        def _handwritten_sub(match: re.Match) -> str:
            fragment = match.group(1).strip()
            name = page_images.get(pages[0]) if pages else None
            if name is not None:
                return f"[HANDWRITTEN: {fragment}](images/{name})"
            return match.group(0)

        block = _IMAGE_PLACEHOLDER_RE.sub(_image_sub, block)
        return _HANDWRITTEN_RE.sub(_handwritten_sub, block)

    def _copy_page_images(self, pdf_path: Path, images_dir: Path) -> dict[int, str]:
        rendered_dir = settings.output_root / "rendered" / pdf_path.stem
        pngs = sorted(rendered_dir.glob("*.png")) if rendered_dir.is_dir() else []
        if not pngs:
            pngs = self._render_pages(pdf_path)
        page_images: dict[int, str] = {}
        for index, png in enumerate(pngs, 1):
            name = f"page_{index:05d}.png"
            _link_or_copy(png, images_dir / name)
            page_images[index] = name
        return page_images

    @staticmethod
    def _render_pages(pdf_path: Path) -> list[Path]:
        try:
            from idp.renderer import extract_pdf_text_and_visual_pages

            _, pngs, _ = extract_pdf_text_and_visual_pages(pdf_path)
            return pngs
        except Exception:
            return []

    def _bundle_docx(self, docx_path: Path, images_dir: Path, text: str) -> str:
        text, counter = _materialize_data_uris(images_dir, text, prefix="docx_image")
        if counter == 0:
            text = self._append_archive_images(text, docx_path, images_dir, member_prefix="word/media/", label="Document")
        return text

    def _bundle_pptx(self, pptx_path: Path, images_dir: Path, text: str) -> str:
        slide_images = _extract_pptx_images(pptx_path, images_dir)
        if not slide_images:
            return self._append_archive_images(text, pptx_path, images_dir, member_prefix="ppt/media/", label="Slide")
        out: list[str] = []
        linked: set[int] = set()
        for line in text.splitlines():
            out.append(line)
            marker = _SLIDE_MARKER_RE.search(line)
            if marker is None:
                continue
            slide_no = int(marker.group(1))
            linked.add(slide_no)
            for name in slide_images.get(slide_no, []):
                out.append(f"![Slide {slide_no} image](images/{name})")
        leftovers = [(slide_no, names) for slide_no, names in sorted(slide_images.items()) if slide_no not in linked]
        if leftovers:
            out.append("")
            out.append("## Slide images")
            for slide_no, names in leftovers:
                for name in names:
                    out.append(f"![Slide {slide_no} image](images/{name})")
        return "\n".join(out)

    def _bundle_html(self, source: Path, images_dir: Path, text: str, *, folder: bool) -> str:
        text, counter = _materialize_data_uris(images_dir, text, prefix="html_image")
        rewritten: list[str] = []
        for base_dirs, section in self._html_sections(source, text, folder=folder):
            section, counter = _rewrite_local_image_refs(section, base_dirs, images_dir, "html_image", counter)
            rewritten.append(section)
        return "\n".join(rewritten)

    @staticmethod
    def _html_sections(source: Path, text: str, *, folder: bool) -> list[tuple[list[Path], str]]:
        if not folder or not source.is_dir():
            base = source.parent if source.is_file() else source
            return [([base], text)]
        sections: list[list[str]] = []
        current: list[str] = []
        for line in text.splitlines():
            if _FILE_MARKER_RE.search(line):
                if current:
                    sections.append(current)
                current = [line]
                continue
            current.append(line)
        if current:
            sections.append(current)
        return [([source], "\n".join(lines)) for lines in sections]

    @staticmethod
    def _append_archive_images(
        text: str,
        archive_path: Path,
        images_dir: Path,
        member_prefix: str,
        label: str,
    ) -> str:
        links: list[str] = []
        try:
            with zipfile.ZipFile(archive_path) as archive:
                members = [
                    name
                    for name in sorted(archive.namelist())
                    if name.startswith(member_prefix) and not name.endswith("/")
                ]
                for index, member in enumerate(members, 1):
                    stored_name = f"{label.lower()}_image_{index:04d}.{_image_extension(Path(member).suffix)}"
                    (images_dir / stored_name).write_bytes(archive.read(member))
                    links.append(f"![{label} image {index}](images/{stored_name})")
        except Exception:
            return text
        if not links:
            return text
        return f"{text.rstrip()}\n\n## {label} images\n\n" + "\n\n".join(links) + "\n"

    def _write_meta(
        self,
        bundle_dir: Path,
        source: Path,
        file_type: str,
        pages: int | None,
        document_type: str,
        annotation: str,
        metadata: dict,
        images_dir: Path,
    ) -> None:
        images = sorted(path.name for path in images_dir.iterdir() if path.is_file())
        meta = {
            "source": self._relative(source),
            "type": file_type,
            "size_bytes": source.stat().st_size if source.is_file() else None,
            "pages": pages,
            "document_type": document_type,
            "annotation": annotation,
            "metadata": metadata,
            "text_file": "text.md",
            "images_dir": "images",
            "images": images,
            "tables_file": "tables.json",
        }
        (bundle_dir / "meta.json").write_text(
            json.dumps(meta, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
