"""Rebuild DOCX, PDF and Visio documents from txts bundles.

Each ``txts/<file name>/`` bundle contains ``text.txt`` (document text with
hyperlinks to images), ``images/`` (page renders and embedded pictures),
``tables.json`` and ``meta.json``. This module parses the text back into
blocks (headings, paragraphs with inline images, tables, lists, page breaks)
and renders it as a real DOCX and/or PDF document with the images in place.

Run standalone:  idp-reconstruct [--bundle DIR] [--root DIR] [--format docx,pdf]
"""

from __future__ import annotations

import json
import logging
import math
import os
import re
from dataclasses import dataclass, field
from pathlib import Path

from idp.config import settings
from idp.visio_document import build_document_vsdx
from idp.visio_exporter import build_bundle_vsdx

logger = logging.getLogger(__name__)

_PAGE_MARKER_RE = re.compile(r"<!--\s*page\s+(\d+)\s*-->")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^[-*+]\s+(.*)$")
_TABLE_ROW_RE = re.compile(r"^\|.*\|$")
_TABLE_SEPARATOR_RE = re.compile(r"^[:\-\s]+$")
_INLINE_RE = re.compile(r"!\[([^\]]*)\]\(([^)\s]+)\)|\[HANDWRITTEN:\s*([^\]]*)\]\(([^)\s]+)\)")
_LINK_ANY_RE = re.compile(r"(!?\[[^\]]*?\])\(([^)]*?)\)", re.DOTALL)
_ESCAPE_RE = re.compile(r"\\([.\-*+\\`()\[\]{}#!|>~])")
_PATH_FRAGMENT_RE = re.compile(r"[A-Za-z0-9_./#%\-]+")

_PAGE_WIDTH = 595.28
_PAGE_HEIGHT = 841.89
_MARGIN = 56.7  # 2 cm
_GAP = 6.0
_FONT_SIZE = 12.0
_TABLE_FONT_SIZE = 11.0
_HEADING_FONT_SIZE = 14.0
_LINE_HEIGHT = 1.5
_MAX_IMAGE_WIDTH_CM = 17.0
_MAX_IMAGE_HEIGHT_CM = 24.0
_FIRST_LINE_INDENT_CM = 1.25
_TEXT_AREA_WIDTH_CM = 17.0
_TEXT_AREA_HEIGHT_CM = 25.7
_DOCX_FONT = "Times New Roman"
_DOCX_FONT_SIZE = 12
_DOCX_HEADING_FONT_SIZE = 14
_DOCX_TABLE_WIDTH_CM = 17.0
_DOCX_TABLE_MIN_COLUMN_CM = 0.9
_FONT_DIR = Path(__file__).parent / "fonts"
_FONT_CANDIDATES = (
    str(_FONT_DIR / "LiberationSerif-Regular.ttf"),
    "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf",
    "/usr/share/fonts/truetype/liberation2/LiberationSerif-Regular.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
)
_MIN_FONT_SIZE = 8.0
_MAX_TEXT_WRITES = 512
_TABLE_MIN_COLUMN_PT = 20.0
_TABLE_HARD_MIN_COLUMN_PT = 9.0
_TABLE_MAX_CELL_WEIGHT = 48
_WRAPPED_LINK_MAX_CHARS = 600


@dataclass
class Segment:
    kind: str
    text: str = ""
    path: str = ""
    alt: str = ""


@dataclass
class Block:
    kind: str
    level: int = 0
    segments: list[Segment] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    items: list[str] = field(default_factory=list)


def _unescape(text: str) -> str:
    return _ESCAPE_RE.sub(r"\1", text)


def _rect_is_finite(rect) -> bool:
    """True when a PyMuPDF rect is finite, non-empty and oriented correctly."""
    try:
        return (
            all(math.isfinite(value) for value in (rect.x0, rect.y0, rect.x1, rect.y1))
            and rect.x1 > rect.x0
            and rect.y1 > rect.y0
        )
    except AttributeError:
        return False


def _split_segments(line: str) -> list[Segment]:
    segments: list[Segment] = []
    position = 0
    for match in _INLINE_RE.finditer(line):
        if match.start() > position:
            segments.append(Segment("text", text=line[position:match.start()]))
        if match.group(2) is not None:
            segments.append(Segment("image", path=match.group(2), alt=match.group(1)))
        else:
            segments.append(Segment("handwritten", text=match.group(3) or "", path=match.group(4) or ""))
        position = match.end()
    if position < len(line):
        segments.append(Segment("text", text=line[position:]))
    return segments


def _join_wrapped_links(text: str) -> str:
    """Repair image links that were wrapped across lines.

    A soft wrap inside a markdown link (alt text or path) breaks the link
    entirely, so the newline is removed and the link reassembled. Only genuine
    image/handwritten links that actually contain a line break are repaired:
    a loose ``](``...``)`` pair far apart (unbalanced brackets in the text)
    must never swallow whole paragraphs, headings or tables between them.
    """

    def _fix(match: re.Match) -> str:
        label = match.group(1)
        path = match.group(2)
        if "\n" not in label and "\n" not in path:
            return match.group(0)  # intact link: never touch it
        if len(label) + len(path) > _WRAPPED_LINK_MAX_CHARS:
            return match.group(0)
        head = label[1:] if label.startswith("!") else label
        if not (head.startswith("[Image:") or head.startswith("[HANDWRITTEN:")):
            return match.group(0)
        joined_path = "".join(path.split())
        if "(" in path or not _PATH_FRAGMENT_RE.fullmatch(joined_path):
            # prose between "]" and ")": an unbalanced bracket, not a link
            return match.group(0)
        label = " ".join(label.split())
        return f"{label}({joined_path})"

    return _LINK_ANY_RE.sub(_fix, text)


def _is_plain_line(line: str) -> bool:
    """A markdown paragraph line (soft wraps are joined into one paragraph)."""
    stripped = line.strip()
    if not stripped:
        return False
    if stripped.startswith("<!--") or stripped.startswith("#") or stripped.startswith("```"):
        return False
    if _TABLE_ROW_RE.match(stripped) and stripped.count("|") >= 2:
        return False
    return _BULLET_RE.match(stripped) is None


def parse_bundle_text(text: str) -> list[Block]:
    """Parse the markdown of a bundle into blocks (wrapped links repaired)."""
    return parse_bundle(_join_wrapped_links(text))


def parse_bundle(text: str) -> list[Block]:
    blocks: list[Block] = []
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("<!--"):
            marker = _PAGE_MARKER_RE.search(stripped)
            if marker is not None:
                blocks.append(Block("page_break", level=int(marker.group(1))))
            index += 1
            continue
        if _TABLE_ROW_RE.match(stripped) and stripped.count("|") >= 2:
            rows: list[list[str]] = []
            while index < len(lines):
                row_line = lines[index].strip()
                if not (_TABLE_ROW_RE.match(row_line) and row_line.count("|") >= 2):
                    break
                cells = [cell.strip() for cell in row_line.strip("|").split("|")]
                if not all(_TABLE_SEPARATOR_RE.fullmatch(cell) for cell in cells if cell):
                    rows.append(cells)
                index += 1
            if rows:
                blocks.append(Block("table", rows=rows))
            continue
        heading = _HEADING_RE.match(stripped)
        if heading is not None:
            blocks.append(Block("heading", level=len(heading.group(1)), segments=[Segment("text", text=heading.group(2).strip())]))
            index += 1
            continue
        if _BULLET_RE.match(stripped):
            items: list[str] = []
            while index < len(lines):
                item = _BULLET_RE.match(lines[index].strip())
                if item is None:
                    break
                items.append(item.group(1).strip())
                index += 1
            blocks.append(Block("list", items=items))
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines) and _is_plain_line(lines[index]):
            paragraph_lines.append(lines[index].strip())
            index += 1
        paragraph = " ".join(part for part in paragraph_lines if part)
        segments = _split_segments(paragraph)
        if segments:
            blocks.append(Block("paragraph", segments=segments))
    return blocks


def _image_size_cm(image_path: Path, bbox: list[float] | None = None) -> tuple[float, float] | None:
    """Display size in cm: from the region bbox when known, else natural 150 dpi."""
    try:
        from PIL import Image

        with Image.open(image_path) as image:
            width_px, height_px = image.size
    except Exception:
        return None
    if not width_px or not height_px:
        return None
    if bbox and len(bbox) == 4 and bbox[2] > bbox[0] and bbox[3] > bbox[1]:
        width_cm = (bbox[2] - bbox[0]) * _TEXT_AREA_WIDTH_CM
        height_cm = (bbox[3] - bbox[1]) * _TEXT_AREA_HEIGHT_CM
        pixel_ratio = width_px / float(height_px) if height_px else 1.0
        target_ratio = width_cm / height_cm if height_cm else pixel_ratio
        if pixel_ratio > target_ratio:
            height_cm = width_cm / pixel_ratio
        else:
            width_cm = height_cm * pixel_ratio
        return width_cm, height_cm
    width_cm = width_px / 150.0 * 2.54
    height_cm = height_px / 150.0 * 2.54
    scale = min(_MAX_IMAGE_WIDTH_CM / width_cm, _MAX_IMAGE_HEIGHT_CM / height_cm, 1.0)
    return width_cm * scale, height_cm * scale


def _table_column_weights(rows: list[list[str]]) -> list[int]:
    """Content weight per column: longest cell, capped so one huge cell
    cannot squeeze every other column into nothing."""
    columns = max((len(row) for row in rows), default=0)
    weights: list[int] = []
    for column in range(columns):
        longest = max(
            (len(row[column]) for row in rows if column < len(row)),
            default=0,
        )
        weights.append(max(4, min(_TABLE_MAX_CELL_WEIGHT, longest)))
    return weights


def _scale_widths(weights: list[int], available: float, minimum: float, hard_minimum: float) -> list[float]:
    """Distribute ``available`` points across ``weights`` with a readable
    minimum width; shrinks proportionally when the minimum does not fit."""
    total = sum(weights)
    if total <= 0 or available <= 0:
        return [available] * len(weights)
    widths = [weight / total * available for weight in weights]
    widths = [max(minimum, width) for width in widths]
    if sum(widths) <= available:
        return widths
    scale = available / sum(widths)
    widths = [max(hard_minimum, width * scale) for width in widths]
    if sum(widths) > available:
        # pathological column count: keep the hard minimum for every column
        # (the table runs past the page edge instead of collapsing/crashing)
        return [hard_minimum] * len(weights)
    return widths


def _table_column_widths(rows: list[list[str]], available: float) -> list[float]:
    """Auto-fitted PDF column widths in points."""
    return _scale_widths(
        _table_column_weights(rows), available, _TABLE_MIN_COLUMN_PT, _TABLE_HARD_MIN_COLUMN_PT
    )


def _docx_table_column_cm(rows: list[list[str]]) -> list[float]:
    """Auto-fitted DOCX column widths in cm."""
    return _scale_widths(
        _table_column_weights(rows), _DOCX_TABLE_WIDTH_CM, _DOCX_TABLE_MIN_COLUMN_CM, 0.5
    )


def _add_docx_table(document, rows: list[list[str]]) -> None:
    from docx.shared import Cm, Pt

    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.autofit = False
    widths = _docx_table_column_cm(rows)
    for column, width in enumerate(widths):
        table.columns[column].width = Cm(width)
    for row_index, row in enumerate(rows):
        for column in range(columns):
            cell = table.cell(row_index, column)
            cell.width = Cm(widths[column])
            cell.text = _unescape(row[column]) if column < len(row) else ""
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = _DOCX_FONT
                    run.font.size = Pt(_DOCX_FONT_SIZE)
                    if row_index == 0 and len(rows) > 1:
                        run.bold = True


def _setup_docx_styles(document) -> None:
    """Times New Roman 12, 1.5 line spacing, 2 cm margins, headings at 14 pt."""
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    for section in document.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    normal = document.styles["Normal"]
    normal.font.name = _DOCX_FONT
    normal.font.size = Pt(_DOCX_FONT_SIZE)
    normal.paragraph_format.line_spacing = _LINE_HEIGHT
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(_FIRST_LINE_INDENT_CM)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), _DOCX_FONT)

    for level in range(1, 7):
        style = document.styles[f"Heading {level}"]
        style.font.name = _DOCX_FONT
        style.font.size = Pt(_DOCX_HEADING_FONT_SIZE)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.line_spacing = _LINE_HEIGHT
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.get_or_add_rFonts()
        for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attribute), _DOCX_FONT)


def build_docx(blocks: list[Block], bundle_dir: Path, output_path: Path, regions: dict[str, dict] | None = None) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    regions = regions or {}
    document = Document()
    _setup_docx_styles(document)
    has_content = False
    for block in blocks:
        if block.kind == "page_break":
            if has_content:
                document.add_page_break()
                has_content = False
        elif block.kind == "heading":
            text = _unescape(block.segments[0].text) if block.segments else ""
            heading = document.add_heading("", level=min(max(block.level, 1), 6))
            run = heading.add_run(text)
            run.font.name = _DOCX_FONT
            run.font.size = Pt(_DOCX_HEADING_FONT_SIZE)
            run.font.bold = True
            has_content = True
        elif block.kind == "list":
            for item in block.items:
                entry = document.add_paragraph(style="List Bullet")
                entry.add_run(_unescape(item))
                entry.paragraph_format.first_line_indent = Cm(0)
                entry.paragraph_format.left_indent = Cm(1.25)
            has_content = True
        elif block.kind == "table":
            _add_docx_table(document, block.rows)
            has_content = True
        elif block.kind == "paragraph":
            only_images = all(segment.kind == "image" for segment in block.segments)
            paragraph = document.add_paragraph()
            if only_images:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Cm(0)
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                paragraph.paragraph_format.first_line_indent = Cm(_FIRST_LINE_INDENT_CM)
            for segment in block.segments:
                if segment.kind == "text":
                    paragraph.add_run(_unescape(segment.text))
                elif segment.kind == "handwritten":
                    image_path = bundle_dir / segment.path
                    if image_path.is_file():
                        run = paragraph.add_run()
                        size = _image_size_cm(image_path, regions.get(Path(segment.path).name, {}).get("bbox"))
                        if size is not None:
                            run.add_picture(str(image_path), width=Cm(size[0]), height=Cm(size[1]))
                        else:
                            run.add_picture(str(image_path))
                        if segment.text:
                            caption = paragraph.add_run(_unescape(segment.text))
                            caption.italic = True
                    else:
                        run = paragraph.add_run(_unescape(segment.text))
                        run.italic = True
                else:
                    image_path = bundle_dir / segment.path
                    if not image_path.is_file():
                        if segment.alt:
                            paragraph.add_run(f"[{segment.alt}]")
                        continue
                    run = paragraph.add_run()
                    size = _image_size_cm(image_path, regions.get(Path(segment.path).name, {}).get("bbox"))
                    if size is not None:
                        run.add_picture(str(image_path), width=Cm(size[0]), height=Cm(size[1]))
                    else:
                        run.add_picture(str(image_path))
            has_content = True
    document.save(str(output_path))
    return output_path


def _resolve_font(bold: bool = False) -> tuple[str, str | None]:
    for candidate in _FONT_CANDIDATES:
        path = Path(candidate)
        if not path.is_file():
            continue
        if not bold:
            return "F0", candidate
        bold_path = path.with_name(path.name.replace("Regular", "Bold"))
        if bold_path.is_file():
            return "FB", str(bold_path)
        return "F0", candidate
    return "china-s", None


def _chars_per_line(width: float, font_size: float) -> int:
    return max(4, int(width / (font_size * 0.52)))


def _split_at_separator(text: str) -> tuple[str, str, str]:
    """Split near the middle at a whitespace boundary so both halves stay wrappable.

    Splitting in the middle keeps the number of retries logarithmic even for a
    paragraph that is far longer than one page.
    """
    target = len(text) // 2
    best: re.Match | None = None
    for match in re.finditer(r"\s+", text):
        if best is None or abs(match.start() - target) < abs(best.start() - target):
            best = match
    if best is None:
        return "", "", text
    return text[: best.start()], text[best.start(): best.end()], text[best.end():]


def _text_lines_needed(text: str, width: float, font_size: float) -> int:
    lines = 0
    for chunk in text.split("\n"):
        lines += max(1, math.ceil(len(chunk) / _chars_per_line(width, font_size)))
    return max(1, lines)


def build_pdf(blocks: list[Block], bundle_dir: Path, output_path: Path, regions: dict[str, dict] | None = None) -> Path:
    import fitz

    regions = regions or {}
    fontname, fontfile = _resolve_font()
    document = fitz.open()
    page = document.new_page(width=_PAGE_WIDTH, height=_PAGE_HEIGHT)
    state = {"y": _MARGIN, "used": False}

    def start_page() -> None:
        nonlocal page
        page = document.new_page(width=_PAGE_WIDTH, height=_PAGE_HEIGHT)
        state["y"] = _MARGIN
        state["used"] = False

    def write_text(text: str, font_size: float = _FONT_SIZE, bold: bool = False,
                   indent: bool = True, align: int = 3) -> None:
        if not text.strip():
            return
        active_font, active_file = _resolve_font(bold)
        body = text
        if indent and align == 3:
            indent_pt = _FIRST_LINE_INDENT_CM * 28.3465
            try:
                font = fitz.Font(fontfile=active_file) if active_file else fitz.Font(active_font)
                space_pt = font.text_length(" ", fontsize=font_size)
            except Exception:
                space_pt = 0.0
            if space_pt > 0:
                body = " " * max(1, int(indent_pt / space_pt)) + text
        pending = [body]
        size = font_size
        writes = 0
        while pending:
            writes += 1
            if writes > _MAX_TEXT_WRITES:
                logger.warning(
                    "PDF reconstruction stopped after %d writes; %d fragment(s) dropped",
                    _MAX_TEXT_WRITES, len(pending),
                )
                break
            remaining = pending.pop(0)
            if not remaining.strip():
                continue
            if _PAGE_HEIGHT - _MARGIN - state["y"] < _MIN_FONT_SIZE * _LINE_HEIGHT:
                start_page()
            rect = fitz.Rect(_MARGIN, state["y"], _PAGE_WIDTH - _MARGIN, _PAGE_HEIGHT - _MARGIN)
            try:
                leftover = page.insert_textbox(
                    rect,
                    remaining,
                    fontsize=size,
                    fontname=active_font,
                    fontfile=active_file,
                    lineheight=_LINE_HEIGHT,
                    align=align,
                )
            except ValueError as exc:
                logger.warning("PDF reconstruction insert_textbox failed: %s", exc)
                start_page()
                pending.insert(0, remaining)
                continue
            if leftover >= 0:
                state["y"] = _PAGE_HEIGHT - _MARGIN - leftover + _GAP
                state["used"] = True
                continue
            if size > _MIN_FONT_SIZE:
                start_page()
                size = max(_MIN_FONT_SIZE, size - 1.0)
                pending.insert(0, remaining)
                continue
            head, separator, tail = _split_at_separator(remaining)
            if not head or not tail:
                logger.warning(
                    "PDF reconstruction dropped a text fragment that does not fit a page: %.120r",
                    remaining,
                )
                start_page()
                state["used"] = False
                continue
            pending.insert(0, tail)
            pending.insert(0, head)

    def write_image(image_path: Path, geometry: dict | None = None) -> None:
        if not image_path.is_file():
            return
        try:
            pixmap = fitz.Pixmap(str(image_path))
        except Exception:
            return
        content_width = _PAGE_WIDTH - 2 * _MARGIN
        content_height = _PAGE_HEIGHT - 2 * _MARGIN
        bbox = (geometry or {}).get("bbox")
        if bbox and len(bbox) == 4 and bbox[2] > bbox[0] and bbox[3] > bbox[1]:
            width = min((bbox[2] - bbox[0]) * content_width, content_width)
            ratio = pixmap.width / float(pixmap.height) if pixmap.height else 1.0
            height = width / ratio if ratio else width
            if height > content_height:
                height = content_height
                width = height * ratio
        else:
            remaining = _PAGE_HEIGHT - _MARGIN - state["y"]
            scale = min(content_width / pixmap.width, remaining / pixmap.height)
            if scale < min(content_width / pixmap.width, content_height / pixmap.height):
                start_page()
                scale = min(content_width / pixmap.width, content_height / pixmap.height)
            width = pixmap.width * scale
            height = pixmap.height * scale
        rect = fitz.Rect(_MARGIN, state["y"], _MARGIN + width, state["y"] + height)
        if not _rect_is_finite(rect):
            logger.warning("PDF reconstruction skipped an image with degenerate geometry: %s", image_path)
            return
        page.insert_image(rect, filename=str(image_path))
        state["y"] += height + _GAP
        state["used"] = True

    def write_table(rows: list[list[str]]) -> None:
        columns = max(len(row) for row in rows)
        available = _PAGE_WIDTH - 2 * _MARGIN
        column_widths = _table_column_widths(rows, available)

        def row_height(row: list[str]) -> float:
            lines = 1
            for column, cell in enumerate(row):
                if not cell:
                    continue
                lines = max(lines, _text_lines_needed(cell, column_widths[column], _TABLE_FONT_SIZE))
            return lines * _TABLE_FONT_SIZE * 1.35 + 6.0

        def draw_row(row: list[str]) -> None:
            height = row_height(row)
            x = _MARGIN
            for column in range(columns):
                cell_rect = fitz.Rect(x, state["y"], x + column_widths[column], state["y"] + height)
                if _rect_is_finite(cell_rect):
                    page.draw_rect(cell_rect, color=(0.55, 0.55, 0.55), width=0.5)
                cell_text = _unescape(row[column]) if column < len(row) else ""
                if cell_text:
                    inner = fitz.Rect(cell_rect.x0 + 2, cell_rect.y0 + 2, cell_rect.x1 - 2, cell_rect.y1 - 2)
                    if _rect_is_finite(inner) and inner.width > 1 and inner.height > 1:
                        try:
                            page.insert_textbox(
                                inner,
                                cell_text,
                                fontsize=_TABLE_FONT_SIZE,
                                fontname=fontname,
                                fontfile=fontfile,
                                lineheight=1.2,
                            )
                        except ValueError as exc:
                            logger.warning(
                                "PDF reconstruction skipped a table cell: %s (%.80r)", exc, cell_text
                            )
                x += column_widths[column]
            state["y"] += height

        header = rows[0] if len(rows) > 1 else None
        index = 0
        while index < len(rows):
            height = row_height(rows[index])
            if state["used"] and state["y"] + height > _PAGE_HEIGHT - _MARGIN:
                start_page()
                if header is not None and index > 0:
                    draw_row(header)
            draw_row(rows[index])
            index += 1
        state["y"] += _GAP
        state["used"] = True

    for block in blocks:
        if block.kind == "page_break":
            if state["used"]:
                start_page()
        elif block.kind == "heading":
            text = _unescape(block.segments[0].text) if block.segments else ""
            write_text(text, _HEADING_FONT_SIZE, bold=True, indent=False, align=0)
        elif block.kind == "list":
            for item in block.items:
                write_text(f"• {_unescape(item)}", indent=False, align=0)
        elif block.kind == "table":
            write_table(block.rows)
        elif block.kind == "paragraph":
            pending = ""
            for segment in block.segments:
                if segment.kind == "image":
                    write_text(pending)
                    pending = ""
                    write_image(bundle_dir / segment.path, regions.get(Path(segment.path).name))
                elif segment.kind == "handwritten":
                    image_path = bundle_dir / segment.path
                    if image_path.is_file():
                        write_text(pending)
                        pending = ""
                        write_image(image_path, regions.get(Path(segment.path).name))
                        if segment.text:
                            write_text(_unescape(segment.text))
                    else:
                        pending += _unescape(segment.text)
                else:
                    pending += _unescape(segment.text)
            write_text(pending)

    document.save(str(output_path))
    return output_path


def load_regions(bundle_dir: Path) -> dict[str, dict]:
    """Image geometry (bbox, pixel size) recorded by the txt exporter."""
    meta_file = bundle_dir / "meta.json"
    if not meta_file.is_file():
        return {}
    try:
        meta = json.loads(meta_file.read_text(encoding="utf-8"))
    except Exception:
        return {}
    regions = meta.get("regions") or []
    return {
        Path(str(region["name"])).name: region
        for region in regions
        if isinstance(region, dict) and region.get("name")
    }


def build_visio(bundle_dir: Path, blocks: list[Block]) -> Path:
    """Build the Visio drawing of a bundle.

    ``IDP_VISIO_CONTENT`` selects the content of the drawing: ``pages`` (the
    page renders, one Visio page per document page), ``sheet`` (all renders on
    a single page) or ``document`` (the editable text of the document).
    """
    bundle_dir = Path(bundle_dir)
    target = bundle_dir / "reconstruct.vsdx"
    content = (settings.visio_content or "pages").lower()
    if content == "document":
        return build_document_vsdx(
            blocks,
            target,
            title=bundle_dir.name,
            base_dir=bundle_dir,
        )
    return build_bundle_vsdx(
        bundle_dir,
        target,
        mode="sheet" if content == "sheet" else "pages",
        title=bundle_dir.name,
    )


def reconstruct_bundle(bundle_dir: Path, formats: list[str] | None = None) -> list[Path]:
    text_file = bundle_dir / "text.txt"
    if not text_file.is_file():
        return []
    requested = formats if formats is not None else list(settings.reconstruct_formats)
    blocks = parse_bundle_text(text_file.read_text(encoding="utf-8"))
    regions = load_regions(bundle_dir)
    created: list[Path] = []
    for fmt in requested:
        fmt = fmt.lower().lstrip(".")
        if fmt == "docx":
            created.append(build_docx(blocks, bundle_dir, bundle_dir / "reconstruct.docx", regions))
        elif fmt == "pdf":
            created.append(build_pdf(blocks, bundle_dir, bundle_dir / "reconstruct.pdf", regions))
        elif fmt == "vsdx":
            try:
                created.append(build_visio(bundle_dir, blocks))
            except ValueError as exc:
                logger.warning("Visio export skipped for %s: %s", bundle_dir, exc)
    return created


def main() -> None:
    import argparse

    logging.basicConfig(
        level=getattr(logging, os.getenv("IDP_LOG_LEVEL", "INFO").upper(), logging.INFO),
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    )
    parser = argparse.ArgumentParser(description="Rebuild DOCX/PDF/Visio documents from txts bundles")
    parser.add_argument("--bundle", type=Path, help="single bundle directory")
    parser.add_argument("--root", type=Path, help="bundles root (default: <output_root>/txts)")
    parser.add_argument("--format", default="", help="comma separated formats: docx,pdf,vsdx")
    args = parser.parse_args()

    formats = [item.strip() for item in args.format.split(",") if item.strip()] or None
    if args.bundle is not None:
        targets = [args.bundle]
    else:
        root = args.root or settings.output_root / "txts"
        targets = sorted(path for path in root.iterdir() if path.is_dir()) if root.is_dir() else []

    for target in targets:
        for created in reconstruct_bundle(target, formats):
            print(f"created {created}")
    if not targets:
        print(f"no bundles found in {args.root or settings.output_root / 'txts'}")


if __name__ == "__main__":
    main()
