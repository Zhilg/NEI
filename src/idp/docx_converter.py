"""Convert DOCX files to Markdown with structure, tables, and noise removal."""

from __future__ import annotations

import re
from pathlib import Path

import mammoth
from docx import Document


_NOISE_PATTERNS = [
    re.compile(r"^\s*(?:Стр\.|Стр\s*:|Page\s+\d+|Page\s*:.*)$", re.IGNORECASE),
    re.compile(r"^\s*(?:Содержание|Оглавление|Contents)\s*$", re.IGNORECASE),
    re.compile(r"^\s*(?:Список\s+литературы|Литература|References|Bibliography)\s*$", re.IGNORECASE),
    re.compile(r"^\s*(?:Благодарности|Acknowledgements?)\s*$", re.IGNORECASE),
    re.compile(r"^\s*(?:Приложения|Appendix)\s*$", re.IGNORECASE),
]


def _clean_noise(text: str) -> str:
    lines = text.splitlines()
    cleaned: list[str] = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            cleaned.append("")
            continue
        if any(p.match(stripped) for p in _NOISE_PATTERNS):
            continue
        cleaned.append(line)
    text = "\n".join(cleaned)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _table_to_markdown(table) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            text = " ".join(
                p.text.strip() for p in cell.paragraphs if p.text.strip()
            )
            cells.append(text)
        rows.append(cells)
    if not rows:
        return ""
    cols = max(len(r) for r in rows)
    normalized = [r + [""] * (cols - len(r)) for r in rows]
    md_lines = []
    for i, row in enumerate(normalized):
        md_lines.append("| " + " | ".join(row) + " |")
        if i == 0:
            md_lines.append("| " + " | ".join(["---"] * cols) + " |")
    return "\n".join(md_lines)


def convert_docx_to_markdown(docx_path: Path) -> str:
    with open(docx_path, "rb") as f:
        result = mammoth.convert_to_markdown(f)
    text = _clean_noise(result.value)
    try:
        doc = Document(str(docx_path))
        table_parts: list[str] = []
        for table in doc.tables:
            table_parts.append(_table_to_markdown(table))
        if table_parts:
            text = text + "\n\n" + "\n\n".join(table_parts)
    except Exception:
        pass
    return text.strip()
