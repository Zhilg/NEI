"""Extract lightweight document metadata from native formats."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any


def extract_docx_metadata(docx_path: Path) -> dict[str, Any]:
    try:
        from docx import Document
        doc = Document(str(docx_path))
        props = doc.core_properties
        data: dict[str, Any] = {}
        if props.title:
            data["title"] = props.title.strip()
        if props.author:
            data["author"] = props.author.strip()
        if props.created:
            data["created"] = props.created.isoformat()
        if props.modified:
            data["modified"] = props.modified.isoformat()
        if props.subject:
            data["subject"] = props.subject.strip()
        if props.keywords:
            data["keywords"] = props.keywords.strip()
        return data
    except Exception:
        return {}


def extract_pptx_metadata(pptx_path: Path) -> dict[str, Any]:
    try:
        from pptx import Presentation
        prs = Presentation(str(pptx_path))
        props = prs.core_properties
        data: dict[str, Any] = {}
        if props.title:
            data["title"] = props.title.strip()
        if props.author:
            data["author"] = props.author.strip()
        if props.created:
            data["created"] = props.created.isoformat()
        if props.modified:
            data["modified"] = props.modified.isoformat()
        if props.subject:
            data["subject"] = props.subject.strip()
        data["slide_count"] = len(prs.slides)
        return data
    except Exception:
        return {}


def extract_metadata(path: Path) -> dict[str, Any]:
    ext = path.suffix.lower()
    if ext == ".docx":
        return extract_docx_metadata(path)
    if ext == ".pptx":
        return extract_pptx_metadata(path)
    return {}
